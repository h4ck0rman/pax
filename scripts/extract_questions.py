"""Local, deterministic question candidate extraction. No OCR or LLM calls.

Output is a staging bank: structural checks do not establish medical accuracy.
"""
from __future__ import annotations
import argparse
from collections import Counter
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
import sys
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools' / 'extraction'))
VERSION = '1'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
SUPPORTED = {'.pdf', '.docx', '.pptx', '.txt', '.json', '.xlsx', '.odt'}
Q = re.compile(r'^\s*(?:(?:Question|Q)\s*)?(\d{1,3})\s*[.):]\s*(.*)$', re.I)
Q_ONLY = re.compile(r'^\s*(?:Question|Q)\s*(\d{1,3})\s*[:.)]?\s*$', re.I)
OPT = re.compile(r'^\s*\(?([A-Ea-e])[.)：:]\s*(.*)$')
ANSWER = re.compile(r'^\s*(?:correct\s+answer|answer|explanation|rationale|reference|discussion)\b', re.I)
VISUAL = re.compile(r'\b(?:shown|pictured|illustrated)\s+(?:below|above)|\b(?:following|this|the)\s+(?:image|figure|diagram|photograph|radiograph|ECG|tracing)\b', re.I)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def block(text, location, **extra):
    return dict(text=text.strip(), location=location, **extra)


def docx_blocks(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read('word/document.xml'))
        definitions, nums, styles, counters = {}, {}, {}, {}
        if 'word/numbering.xml' in z.namelist():
            nr = ET.fromstring(z.read('word/numbering.xml'))
            for abstract in nr.findall(W + 'abstractNum'):
                definitions[abstract.get(W + 'abstractNumId')] = {
                    int(l.get(W + 'ilvl')): l for l in abstract.findall(W + 'lvl')}
            for n in nr.findall(W + 'num'):
                nums[n.get(W + 'numId')] = n
        if 'word/styles.xml' in z.namelist():
            for s in ET.fromstring(z.read('word/styles.xml')).findall(W + 'style'):
                styles[s.get(W + 'styleId')] = s

        def numbering(p):
            properties = p.find(W + 'pPr')
            if properties is None:
                return None
            np = properties.find(W + 'numPr')
            style = properties.find(W + 'pStyle')
            visited = set()
            while np is None and style is not None:
                sid = style.get(W + 'val')
                if sid in visited or sid not in styles:
                    break
                visited.add(sid)
                s = styles[sid]
                np = s.find(W + 'pPr/' + W + 'numPr')
                style = s.find(W + 'basedOn')
            if np is None or np.find(W + 'numId') is None:
                return None
            nid = np.find(W + 'numId').get(W + 'val')
            level = np.find(W + 'ilvl')
            ilvl = int(level.get(W + 'val')) if level is not None else 0
            if nid not in nums:
                return None
            n = nums[nid]
            aid = n.find(W + 'abstractNumId').get(W + 'val')
            lvl = definitions.get(aid, {}).get(ilvl)
            override = next((o for o in n.findall(W + 'lvlOverride') if int(o.get(W + 'ilvl')) == ilvl), None)
            if override is not None and override.find(W + 'lvl') is not None:
                lvl = override.find(W + 'lvl')
            if lvl is None:
                return None
            fmt = lvl.find(W + 'numFmt')
            fmt = fmt.get(W + 'val') if fmt is not None else 'unknown'
            start = lvl.find(W + 'start')
            start = int(start.get(W + 'val')) if start is not None else 1
            if override is not None and override.find(W + 'startOverride') is not None:
                start = int(override.find(W + 'startOverride').get(W + 'val'))
            key = (nid, ilvl)
            counters[key] = counters.get(key, start - 1) + 1
            for k in list(counters):
                if k[0] == nid and k[1] > ilvl:
                    del counters[k]
            return {'format': fmt, 'value': counters[key], 'level': ilvl, 'list_id': nid}

        result = []
        for index, p in enumerate(root.iter(W + 'p'), 1):
            text = ''.join((t.text or '') if t.tag == W + 't' else '\t' if t.tag == W + 'tab' else '\n'
                           for t in p.iter() if t.tag in (W + 't', W + 'tab', W + 'br'))
            meta = numbering(p)
            result.append(block(text, f'paragraph:{index}', numbering=meta,
                                image=any(t.tag in (W + 'drawing', W + 'pict') for t in p.iter()),
                                highlighted=p.find('.//' + W + 'highlight') is not None))
        return result, {}


def read_blocks(data, ext):
    if ext == '.docx':
        return docx_blocks(data)
    if ext == '.pdf':
        import pymupdf
        pymupdf.TOOLS.mupdf_display_errors(False)
        pymupdf.TOOLS.mupdf_display_warnings(False)
        result, low, errors = [], [], []
        with pymupdf.open(stream=data, filetype='pdf') as pdf:
            if pdf.needs_pass:
                raise ValueError('Password-protected PDF')
            count = len(pdf)
            for i, page in enumerate(pdf):
                try:
                    text = page.get_text(sort=True)
                    if len(text.strip()) < 80:
                        low.append(i + 1)
                    result.extend(block(line, f'page:{i+1}:line:{j+1}') for j, line in enumerate(text.splitlines()))
                except Exception as exc:
                    errors.append({'page': i + 1, 'error': str(exc)})
        return result, dict(pages=count, low_text_pages=low, page_errors=errors)
    if ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        result = []
        def shape_text(shapes, loc):
            for shape in shapes:
                if shape.shape_type == 6:
                    shape_text(shape.shapes, loc)
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        result.append(block(p.text, loc))
                if shape.has_table:
                    for row in shape.table.rows:
                        result.append(block('\t'.join(c.text for c in row.cells), loc))
        for i, slide in enumerate(prs.slides, 1):
            shape_text(slide.shapes, f'slide:{i}')
            # Notes remain separate: never append them to a final answer option.
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                result.append(block('Explanation: ' + slide.notes_slide.notes_text_frame.text, f'slide:{i}:notes'))
        return result, dict(slides=len(prs.slides))
    if ext == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        result = [block('\t'.join('' if v is None else str(v) for v in row), f'sheet:{s.title}:row:{i}')
                  for s in book for i, row in enumerate(s.iter_rows(values_only=True), 1)]
        book.close()
        return result, {}
    if ext == '.odt':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            r = ET.fromstring(z.read('content.xml'))
            ns = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
            return [block(''.join(p.itertext()), f'paragraph:{i}') for i, p in enumerate(r.iter(ns+'p'), 1)], {}
    encoding = 'utf-8-sig'
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError:
        encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'cp1252'
        text = data.decode(encoding)
    if ext == '.json':
        # Preserve JSON structure for future schema-specific adapters, without guessing keys.
        obj = json.loads(text)
        result = []
        def walk(v, loc):
            if isinstance(v, str):
                result.extend(block(t, loc) for t in v.splitlines())
            elif isinstance(v, list):
                for i, child in enumerate(v):
                    walk(child, f'{loc}/{i}')
            elif isinstance(v, dict):
                for k, child in v.items():
                    walk(child, f'{loc}/{k}')
        walk(obj, '$')
        return result, {'encoding': encoding, 'schema_adapter_needed': True}
    return [block(t, f'line:{i}') for i, t in enumerate(text.splitlines(), 1)], {'encoding': encoding}


def tokens(blocks):
    result = []
    for b in blocks:
        for line in b['text'].splitlines() or ['']:
            t = dict(b, text=line.strip(), kind='text', label=None)
            n = b.get('numbering')
            q = Q.match(t['text']) or Q_ONLY.match(t['text'])
            o = OPT.match(t['text'])
            if o:
                t.update(kind='option', label=o[1].upper(), text=o[2])
            elif q:
                t.update(kind='question', label=q[1], text=q[2] if len(q.groups()) > 1 else '')
            elif n and n['format'] in ('upperLetter', 'lowerLetter') and 1 <= n['value'] <= 5:
                t.update(kind='option', label=chr(64 + n['value']))
            elif n and n['format'] == 'decimal':
                t.update(kind='question', label=str(n['value']))
            elif ANSWER.match(t['text']):
                t['kind'] = 'answer'
            result.append(t)
    return result


def parse_questions(blocks):
    ts = tokens(blocks)
    starts = [i for i, t in enumerate(ts) if t['kind'] == 'option' and t['label'] == 'A']
    candidates = []
    previous_end = 0
    for a in starts:
        # Limit inferred stems to nearby text and prefer explicit question boundaries.
        lower = max(previous_end, a - 100)
        boundaries = [i for i in range(lower, a) if ts[i]['kind'] in ('question', 'answer', 'option')]
        qstart = boundaries[-1] if boundaries else lower
        explicit = bool(boundaries and ts[qstart]['kind'] == 'question')
        # A numbered list inside a stem is not a sequence of new questions.
        # Word provides list IDs, allowing us to distinguish the nested list.
        if explicit and ts[qstart].get('numbering'):
            last_list = ts[qstart]['numbering']['list_id']
            for k in reversed(boundaries[:-1]):
                prior = ts[k]
                if prior['kind'] != 'question':
                    break
                if (prior.get('numbering') and prior['numbering']['list_id'] != last_list
                        and re.search(r'\?|\b(?:which|what)\b', prior['text'], re.I)):
                    qstart = k
                    break
        if not explicit and boundaries:
            qstart += 1
        stem_tokens = ts[qstart:a]
        stem = '\n'.join((str(t['label']) + '. ' if i > 0 and t['kind'] == 'question' else '') + t['text']
                         for i, t in enumerate(stem_tokens) if t['text']).strip()
        opts, end, tail_trimmed = [], a, False
        j = a
        while j < len(ts):
            t = ts[j]
            if t['kind'] != 'option' or t['label'] != chr(65 + len(opts)):
                break
            label, parts = t['label'], [t['text']]
            j += 1
            while j < len(ts) and ts[j]['kind'] == 'text':
                # Paragraph breaks after an option can indicate a new unnumbered stem.
                if not ts[j]['text']:
                    look = j + 1
                    while look < len(ts) and not ts[look]['text']:
                        look += 1
                    if look < len(ts) and ts[look]['kind'] == 'text' and len(opts) >= 3:
                        tail_trimmed = True
                        break
                parts.append(ts[j]['text'])
                j += 1
                if sum(map(len, parts)) > 2500:
                    tail_trimmed = True
                    break
            opts.append({'label': label, 'text': '\n'.join(p for p in parts if p).strip()})
            end = j
            if tail_trimmed or len(opts) == 5:
                break
        if len(opts) < 2 or len(stem) < 20:
            continue
        flags = []
        if not explicit:
            flags.append('inferred_stem_boundary')
        if len(opts) not in (4, 5):
            flags.append('unusual_option_count')
        if any(not o['text'] for o in opts):
            flags.append('empty_option')
        if any(len(o['text']) > 700 for o in opts):
            flags.append('long_option_possible_contamination')
        if not re.search(r'\?|\b(?:which|what|most likely|best|except|true|false)\b', stem, re.I):
            flags.append('no_question_cue')
        if any(re.search(r'https?://|\bet al\b|\b(?:correct|incorrect)\s*[-:]', o['text'], re.I) for o in opts):
            flags.append('possible_explanation_or_reference_in_options')
        if len(stem) > 4000:
            flags.append('long_stem_possible_contamination')
        if tail_trimmed:
            flags.append('uncertain_option_end')
        raw = ts[qstart:end]
        if VISUAL.search(stem) or any(t.get('image') for t in raw):
            flags.append('visual_dependency')
        if '\ufffd' in stem or any('\ufffd' in o['text'] for o in opts):
            flags.append('encoding_damage')
        if end < len(ts) and ts[end]['kind'] == 'option':
            flags.append('unexpected_option_sequence')
        if any(t.get('numbering') for t in raw):
            # Numbering formats are read, but exotic Word restart rules need review.
            flags.append('word_numbering_reconstructed')
        text_key = re.sub(r'\s+', ' ', json.dumps([stem, opts], ensure_ascii=False)).casefold()
        candidates.append(dict(id=digest(text_key.encode()), number=ts[qstart]['label'] if explicit else None,
                               stem=stem, options=opts, flags=sorted(set(flags)),
                               structural_status='needs_review' if flags else 'structurally_clean',
                               start=ts[qstart]['location'], end=ts[max(qstart, end-1)]['location'],
                               raw_text='\n'.join(t['text'] for t in raw)))
        previous_end = end
    return candidates


def run(resources, output):
    output.mkdir(parents=True, exist_ok=True)
    cache = output / 'documents'
    cache.mkdir(exist_ok=True)
    rows, occurrences, unique, seen = [], [], {}, {}
    counters = Counter()

    def process(path, data, ext, depth=0):
        row = dict(path=path, extension=ext, bytes=len(data), status='', reason='', candidates=0,
                   structurally_clean=0, needs_review=0, duplicate_of='', sha256=digest(data))
        rows.append(row)
        member_name = Path(path.split('!/')[-1])
        if '__MACOSX' in member_name.parts or member_name.name.startswith('._') or member_name.name == '.DS_Store':
            row.update(status='skipped', reason='macOS metadata sidecar')
            return
        if not data or Path(path.split('!/')[-1]).name.startswith('~$'):
            row.update(status='skipped', reason='Empty file or Office lock file')
            return
        if ext == '.zip':
            if depth >= 3:
                row.update(status='deferred', reason='Archive nesting limit')
                return
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    row['status'] = 'archive_processed'
                    for member_index, info in enumerate(z.infolist(), 1):
                        if info.is_dir():
                            continue
                        member = f'{path}!/{info.filename}'
                        suffix = Path(info.filename).suffix.lower()
                        if info.file_size > 250_000_000:
                            rows.append(dict(path=member, extension=suffix, bytes=info.file_size, status='deferred', reason='Archive member exceeds 250 MB limit', candidates=0))
                        else:
                            try:
                                process(member, z.read(info), suffix, depth + 1)
                            except Exception as exc:
                                rows.append(dict(path=member, extension=suffix, bytes=info.file_size, status='failed', reason=str(exc), candidates=0))
                        if member_index % 50 == 0:
                            print(f'Archive {Path(path).name}: {member_index} members inspected', flush=True)
            except Exception as exc:
                row.update(status='failed', reason=str(exc))
            return
        if ext not in SUPPORTED:
            row.update(status='deferred', reason='Legacy Office conversion unavailable' if ext in ('.doc', '.ppt') else 'Outside text extraction scope')
            return
        counters['supported_file_occurrences'] += 1
        sha = row['sha256']
        if sha in seen:
            original, qs = seen[sha]
            row.update(status='duplicate', duplicate_of=original)
        else:
            try:
                target = cache / f'{sha}.json'
                cached = json.loads(target.read_text(encoding='utf-8')) if target.exists() else None
                if cached and cached.get('version') == VERSION:
                    blocks, meta = cached['blocks'], cached['metadata']
                    counters['cache_hits'] += 1
                else:
                    blocks, meta = read_blocks(data, ext)
                    target.write_text(json.dumps(dict(version=VERSION, blocks=blocks, metadata=meta), ensure_ascii=False), encoding='utf-8')
                    counters['documents_text_extracted'] += 1
                qs = parse_questions(blocks)
                numeric_markers = [t['label'] for t in tokens(blocks) if t['kind'] == 'question']
                row.update(status='extracted' if qs else 'no_questions_detected',
                           characters=sum(len(b['text']) for b in blocks),
                           numbered_markers=len(numeric_markers),
                           unmatched_numbered_markers=max(0, len(numeric_markers)-sum(q['number'] is not None for q in qs)), **meta)
                if meta.get('page_errors'):
                    row.update(status='partial', reason='Some PDF pages failed; see page_errors')
                elif meta.get('low_text_pages'):
                    row.update(status='partial' if qs else 'no_questions_detected', reason='Low-text pages present; OCR deferred (may include blank/diagram pages)')
                seen[sha] = (path, qs)
            except Exception as exc:
                row.update(status='failed', reason=f'{type(exc).__name__}: {exc}')
                return
        row['candidates'] = len(qs)
        row['structurally_clean'] = sum(q['structural_status'] == 'structurally_clean' for q in qs)
        row['needs_review'] = len(qs) - row['structurally_clean']
        for q in qs:
            if q['id'] in unique:
                prior = unique[q['id']]
                prior['flags'] = sorted(set(prior['flags']) | set(q['flags']))
                prior['structural_status'] = 'needs_review' if prior['flags'] else 'structurally_clean'
            else:
                unique[q['id']] = dict(q)
            occurrences.append(dict(question_id=q['id'], source=path, number=q['number'], start=q['start'], end=q['end'], flags=q['flags'], structural_status=q['structural_status']))

    paths = sorted(p for p in resources.rglob('*') if p.is_file())
    for i, p in enumerate(paths, 1):
        relative = p.relative_to(resources).as_posix()
        if p.parent == resources and p.suffix.lower() == '.zip':
            rows.append(dict(path=relative, extension='.zip', bytes=p.stat().st_size, status='skipped', reason='Original export archive; expanded folders processed', candidates=0))
        elif p.suffix.lower() not in SUPPORTED | {'.zip', '.doc', '.ppt'}:
            rows.append(dict(path=relative, extension=p.suffix.lower(), bytes=p.stat().st_size, status='deferred', reason='Outside text extraction scope', candidates=0))
        else:
            try:
                counters['physical_files_read'] += 1
                process(relative, p.read_bytes(), p.suffix.lower())
            except Exception as exc:
                rows.append(dict(path=relative, extension=p.suffix.lower(), bytes=p.stat().st_size, status='failed', reason=str(exc), candidates=0))
        if i % 25 == 0:
            print(f'{i}/{len(paths)} physical files; {len(unique)} unique candidates', flush=True)
    db = sqlite3.connect(output / 'questions.sqlite')
    with db:
        db.executescript('''
        DROP TABLE IF EXISTS options; DROP TABLE IF EXISTS occurrences;
        DROP TABLE IF EXISTS questions; DROP TABLE IF EXISTS documents;
        CREATE TABLE documents(path TEXT PRIMARY KEY, status TEXT, metadata_json TEXT);
        CREATE TABLE questions(id TEXT PRIMARY KEY, stem TEXT, structural_status TEXT, flags_json TEXT, raw_text TEXT, review_status TEXT DEFAULT 'unreviewed');
        CREATE TABLE options(question_id TEXT REFERENCES questions(id), label TEXT, text TEXT, PRIMARY KEY(question_id,label));
        CREATE TABLE occurrences(question_id TEXT REFERENCES questions(id), source TEXT REFERENCES documents(path), original_number TEXT, start TEXT, end TEXT, flags_json TEXT, structural_status TEXT);
        ''')
        db.executemany('INSERT INTO documents VALUES (?,?,?)', [(r['path'], r['status'], json.dumps(r, ensure_ascii=False)) for r in rows])
        for q in unique.values():
            db.execute('INSERT INTO questions(id,stem,structural_status,flags_json,raw_text) VALUES (?,?,?,?,?)',
                       (q['id'], q['stem'], q['structural_status'], json.dumps(q['flags']), q['raw_text']))
            db.executemany('INSERT INTO options VALUES (?,?,?)', [(q['id'], o['label'], o['text']) for o in q['options']])
        db.executemany('INSERT INTO occurrences VALUES (?,?,?,?,?,?,?)',
                       [(o['question_id'], o['source'], o['number'], o['start'], o['end'], json.dumps(o['flags']), o['structural_status']) for o in occurrences])
    integrity = db.execute('PRAGMA integrity_check').fetchone()[0]
    foreign_key_issues = db.execute('PRAGMA foreign_key_check').fetchall()
    db.close()
    summary = dict(physical_files_inventoried=len(paths), total_file_records_including_archive_members=len(rows),
                   archive_members=sum('!/' in r['path'] for r in rows),
                   files_with_candidates=sum(r['candidates'] > 0 for r in rows),
                   physical_files_read=counters['physical_files_read'],
                   unique_documents_with_text=sum(r['status'] in ('extracted', 'partial', 'no_questions_detected') for r in rows),
                   pdf_pages_read=sum(r.get('pages', 0) for r in rows),
                   low_text_pdf_pages=sum(len(r.get('low_text_pages', [])) for r in rows),
                   statuses=dict(Counter(r['status'] for r in rows)),
                   unique_question_candidates=len(unique), question_occurrences=len(occurrences),
                   duplicate_question_occurrences=len(occurrences)-len(unique),
                   unique_structural_statuses=dict(Counter(q['structural_status'] for q in unique.values())),
                   occurrence_flags=dict(Counter(f for o in occurrences for f in o['flags'])),
                   processing=dict(counters), sqlite_integrity=integrity, sqlite_foreign_key_issues=foreign_key_issues,
                   limitations=['All questions are unreviewed candidates, not verified quiz content.',
                                'No OCR, LLM, or answer generation was used.',
                                'Legacy DOC/PPT conversion unavailable; these files are deferred.',
                                'Low-text pages may be blank, graphical, or scanned.',
                                'Question recall and precision have not been measured against a labelled dataset.',
                                'Exact text deduplication only; near duplicates remain.',
                                'JSON strings are flattened; arbitrary JSON schemas need an adapter.'])
    for name, obj in [('files', rows), ('summary', summary)]:
        (output / f'{name}.json').write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
    for name, items in [('questions', unique.values()), ('occurrences', occurrences)]:
        with (output / f'{name}.jsonl').open('w', encoding='utf-8') as f:
            for item in items:
                f.write(json.dumps(item, ensure_ascii=False) + '\n')
    fields = ['path', 'extension', 'bytes', 'status', 'candidates', 'structurally_clean', 'needs_review', 'duplicate_of', 'reason']
    with (output / 'files.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)
    report = ['# Pax question extraction report', '',
              '**Python-only staging extraction. Every candidate remains unreviewed.**', '',
              '## Counts', '', *[f'- {k}: {v}' for k, v in summary.items() if isinstance(v, int)], '',
              '## File statuses', '', *[f'- {k}: {v}' for k,v in summary['statuses'].items()], '',
              '## Candidate structural checks', '', *[f'- {k}: {v}' for k,v in summary['unique_structural_statuses'].items()], '',
              '## Limitations', '', *['- '+s for s in summary['limitations']], '',
              '## Outputs', '', '- `files.csv`: every physical file and inspected archive member, status, counts, and reason.',
              '- `files.json`: detailed page-level issues and source metadata.',
              '- `questions.sqlite`: staging questions, options, documents, and occurrences.',
              '- `questions.jsonl` / `occurrences.jsonl`: readable exports.',
              '- `documents/`: cached structured text with source locations.', '',
              'A file with extracted candidates may contain additional undetected questions. Candidate counts are not exam completeness guarantees.']
    (output / 'report.md').write_text('\n'.join(report)+'\n', encoding='utf-8')
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--resources', type=Path, default=ROOT/'resources')
    ap.add_argument('--output', type=Path, default=ROOT/'data'/'extraction')
    args = ap.parse_args()
    if not args.resources.is_dir() or args.output.resolve().is_relative_to(args.resources.resolve()):
        ap.error('Resources must exist and output must be outside resources')
    run(args.resources, args.output)
